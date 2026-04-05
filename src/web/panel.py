# src/web/panel.py
import logging
from pathlib import Path

from datetime import datetime

from fastapi import APIRouter, Depends, File, Form, HTTPException, Request, Response, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, select, update as sa_update, delete
from sqlalchemy.orm import selectinload

from src.db.database import async_session
from src.db.models import (
    TripEventType,
    Driver, DriverDocument, Expense, FuelAnomaly, LocationUpdate,
    RepairStatus, Trip, TripDocument, TripStatus, Vehicle,
    VehicleDocument, VehicleRepair, VehicleStatus, DriverRequest,
    RequestStatus, Geofence, TripEvent, PanelUser
)
from src.services.r2 import save_uploaded_file
from datetime import date as date_type
from src.web.auth import (
    create_access_token,
    get_current_user_from_cookie,
    hash_password,
    verify_password,
)
from src.core.config import settings
from src.services.fuel_antifraud import FuelAntifraudService
from src.services.maintenance import get_maintenance_dashboard_data

logger = logging.getLogger(__name__)


async def get_pending_requests_count() -> int:
    """Get count of pending driver requests for sidebar badge."""
    async with async_session() as session:
        result = await session.execute(
            select(func.count(DriverRequest.id)).where(DriverRequest.status == RequestStatus.PENDING)
        )
        return result.scalar() or 0


async def notify_driver_expense_decision(expense: Expense, approved: bool, rejection_comment: str = None) -> None:
    """Send Telegram message to driver about expense approval or rejection."""
    from src.main import ptb_app  # deferred import — ptb_app is set at lifespan startup
    if ptb_app is None:
        return
    try:
        driver_tg_id = expense.trip.driver.tg_id
        if approved:
            text = f"Расход одобрен: {expense.expense_type} {expense.amount} {expense.currency}"
        else:
            text = f"Расход отклонен: {expense.expense_type} {expense.amount} {expense.currency}"
            if rejection_comment:
                text += f"\nПричина: {rejection_comment}"
        await ptb_app.bot.send_message(chat_id=driver_tg_id, text=text)
    except Exception as e:
        logger.error(f"Failed to notify driver: {e}")


async def notify_driver_trip_assignment(trip: Trip, open_repairs: list = None) -> None:
    """Send Telegram message to driver about new trip assignment."""
    from src.main import ptb_app  # deferred import — ptb_app is set at lifespan startup
    if ptb_app is None:
        return
    try:
        driver_tg_id = trip.driver.tg_id
        if not driver_tg_id:
            return  # dispatcher-created driver, no Telegram account
        text = (
            f"Вам назначен новый рейс #{trip.id}!\n"
            f"Маршрут: {trip.route_details or '—'}\n"
            f"Груз: {trip.load_info or '—'}\n"
            f"Клиент: {trip.client_name or '—'}"
        )
        if open_repairs:
            priority_labels = {"CRITICAL": "Критический", "HIGH": "Высокий", "NORMAL": "Нормальный", "LOW": "Низкий"}
            repairs_text = "\n".join(
                f"• {r.description} ({priority_labels.get(r.priority, r.priority)})"
                for r in open_repairs
            )
            text += f"\n\n⚠️ Внимание! На борту есть неисправности:\n{repairs_text}"
        await ptb_app.bot.send_message(chat_id=driver_tg_id, text=text)
    except Exception as e:
        logger.error(f"Failed to notify driver about trip assignment: {e}")



router = APIRouter(prefix="/panel", tags=["panel"])
templates = Jinja2Templates(directory=str(Path(__file__).parent / "templates"))


@router.get("/login", response_class=HTMLResponse)
async def login_page(request: Request):
    return templates.TemplateResponse(request, "login.html", {"request": request})


@router.post("/login")
async def login_submit(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
):
    role = None
    db_user = None

    # 1. Check DB panel_users first
    try:
        async with async_session() as session:
            result = await session.execute(
                select(PanelUser).where(PanelUser.username == username)
            )
            db_user = result.scalars().first()
    except Exception as e:
        logger.warning(f"panel_users table not available (migration pending?): {e}")

    if db_user:
        if not db_user.is_active:
            return templates.TemplateResponse(
                request, "login.html",
                {"request": request, "error": "Аккаунт ожидает активации администратором"},
                status_code=401,
            )
        if verify_password(password, db_user.password_hash):
            role = db_user.role

    # 2. Fallback: env var admin
    if role is None and username == settings.ADMIN_USERNAME:
        try:
            if settings.ADMIN_PASSWORD_HASH and verify_password(password, settings.ADMIN_PASSWORD_HASH):
                role = "admin"
        except Exception as e:
            logger.error(f"Admin password hash is invalid: {e}. Set ADMIN_PASSWORD in Railway env vars.")
            role = None

    if role is None:
        return templates.TemplateResponse(
            request, "login.html",
            {"request": request, "error": "Неверный логин или пароль"},
            status_code=401,
        )

    token = create_access_token({"sub": username, "role": role})
    response = RedirectResponse(url="/panel/drivers", status_code=302)
    response.set_cookie(
        "access_token", token,
        httponly=True,
        max_age=settings.JWT_EXPIRE_MINUTES * 60,
        samesite="lax",
    )
    return response


@router.get("/logout")
async def logout():
    response = RedirectResponse(url="/panel/login", status_code=302)
    response.delete_cookie("access_token")
    return response


# --- REGISTRATION ---

@router.get("/register", response_class=HTMLResponse)
async def register_page(request: Request):
    return templates.TemplateResponse(request, "register.html", {"request": request})


@router.post("/register", response_class=HTMLResponse)
async def register_submit(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    confirm_password: str = Form(...),
):
    if password != confirm_password:
        return templates.TemplateResponse(request, "register.html",
            {"request": request, "error": "Пароли не совпадают"}, status_code=400)
    if len(password) < 6:
        return templates.TemplateResponse(request, "register.html",
            {"request": request, "error": "Пароль минимум 6 символов"}, status_code=400)
    async with async_session() as session:
        existing = await session.execute(select(PanelUser).where(PanelUser.username == username))
        if existing.scalars().first():
            return templates.TemplateResponse(request, "register.html",
                {"request": request, "error": "Логин уже занят"}, status_code=400)
        new_user = PanelUser(
            username=username,
            password_hash=hash_password(password),
            role="pending",
            is_active=False,
        )
        session.add(new_user)
        await session.commit()
    return templates.TemplateResponse(request, "register.html",
        {"request": request, "success": True})


# --- USER MANAGEMENT (admin only) ---

@router.get("/users", response_class=HTMLResponse)
async def panel_users_list(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    if current_user.get("role") != "admin":
        return RedirectResponse(url="/panel/drivers", status_code=302)
    async with async_session() as session:
        result = await session.execute(select(PanelUser).order_by(PanelUser.created_at.desc()))
        users = result.scalars().all()
    return templates.TemplateResponse(request, "panel_users.html",
        {"request": request, "user": current_user, "panel_users": users})


@router.post("/users/{user_id}/activate")
async def activate_panel_user(
    user_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    if current_user.get("role") != "admin":
        return RedirectResponse(url="/panel/drivers", status_code=302)
    async with async_session() as session:
        result = await session.execute(select(PanelUser).where(PanelUser.id == user_id))
        u = result.scalars().first()
        if u:
            u.is_active = True
            session.add(u)
            await session.commit()
    return RedirectResponse(url="/panel/users", status_code=302)


@router.post("/users/{user_id}/role")
async def set_panel_user_role(
    user_id: int,
    role: str = Form(...),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    if current_user.get("role") != "admin":
        return RedirectResponse(url="/panel/drivers", status_code=302)
    if role not in ("dispatcher", "director", "admin", "pending"):
        return RedirectResponse(url="/panel/users", status_code=302)
    async with async_session() as session:
        result = await session.execute(select(PanelUser).where(PanelUser.id == user_id))
        u = result.scalars().first()
        if u:
            u.role = role
            session.add(u)
            await session.commit()
    return RedirectResponse(url="/panel/users", status_code=302)


@router.post("/users/{user_id}/delete")
async def delete_panel_user(
    user_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    if current_user.get("role") != "admin":
        return RedirectResponse(url="/panel/drivers", status_code=302)
    async with async_session() as session:
        result = await session.execute(select(PanelUser).where(PanelUser.id == user_id))
        u = result.scalars().first()
        if u:
            await session.delete(u)
            await session.commit()
    return RedirectResponse(url="/panel/users", status_code=302)


@router.get("/drivers", response_class=HTMLResponse)
async def list_drivers(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(Driver, Vehicle)
            .outerjoin(Vehicle, Driver.vehicle_id == Vehicle.id)
            .order_by(Driver.created_at.desc())
        )
        rows = result.all()  # List of (Driver, Vehicle|None) tuples
    prc = await get_pending_requests_count()
    return templates.TemplateResponse(request, "drivers.html",
        {"request": request, "rows": rows, "user": current_user, "pending_requests_count": prc},
    )


@router.post("/drivers/{driver_id}/activate")
async def activate_driver(
    driver_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Driver).where(Driver.id == driver_id))
        driver = result.scalars().first()
        if driver and driver.status != "ACTIVE":
            driver.status = "ACTIVE"
            session.add(driver)
            await session.commit()
    return RedirectResponse(url="/panel/drivers", status_code=302)


@router.post("/drivers/{driver_id}/delete")
async def driver_delete(
    driver_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Driver).where(Driver.id == driver_id))
        driver = result.scalars().first()
        if driver:
            # Clean up LocationUpdate manually (no cascade defined)
            await session.execute(delete(LocationUpdate).where(LocationUpdate.driver_id == driver_id))
            
            # Find trips for this driver
            trip_ids = (await session.execute(select(Trip.id).where(Trip.driver_id == driver_id))).scalars().all()
            if trip_ids:
                # Clean up FuelAnomaly and reset trip_id on LocationUpdate/VehicleRepair
                await session.execute(delete(FuelAnomaly).where(FuelAnomaly.trip_id.in_(trip_ids)))
                await session.execute(sa_update(LocationUpdate).where(LocationUpdate.trip_id.in_(trip_ids)).values(trip_id=None))
                await session.execute(sa_update(VehicleRepair).where(VehicleRepair.trip_id.in_(trip_ids)).values(trip_id=None))
                await session.execute(sa_update(DriverRequest).where(DriverRequest.trip_id.in_(trip_ids)).values(trip_id=None))
            
            await session.delete(driver)
            await session.commit()
    return RedirectResponse(url="/panel/drivers", status_code=302)


@router.post("/vehicles/{vehicle_id}/delete")
async def vehicle_delete(
    vehicle_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
        if vehicle:
            # Find trips for this vehicle
            trip_ids = (await session.execute(select(Trip.id).where(Trip.vehicle_id == vehicle_id))).scalars().all()
            if trip_ids:
                await session.execute(sa_update(LocationUpdate).where(LocationUpdate.trip_id.in_(trip_ids)).values(trip_id=None))
                await session.execute(sa_update(DriverRequest).where(DriverRequest.trip_id.in_(trip_ids)).values(trip_id=None))
                # Note: VehicleRepairs are cascade deleted by Vehicle anyway
            
            # Clean up FuelAnomaly manually
            await session.execute(delete(FuelAnomaly).where(FuelAnomaly.vehicle_id == vehicle_id))

            await session.delete(vehicle)
            await session.commit()
    return RedirectResponse(url="/panel/drivers", status_code=302)


@router.get("/vehicles/{vehicle_id}/edit", response_class=HTMLResponse)
async def vehicle_edit_form(
    request: Request,
    vehicle_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
    if not vehicle:
        return RedirectResponse(url="/panel/drivers", status_code=302)
    return templates.TemplateResponse(request, "vehicle_edit.html",
        {"request": request, "vehicle": vehicle, "user": current_user},
    )


@router.post("/vehicles/{vehicle_id}/edit")
async def vehicle_edit_submit(
    request: Request,
    vehicle_id: int,
    make: str = Form(default=""),
    model: str = Form(default=""),
    year: str = Form(default=""),
    fuel_norm_per_100km: str = Form(default=""),
    maintenance_interval_km: str = Form(default=""),
    last_maintenance_odometer: str = Form(default=""),
    trailer_plate: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
        if vehicle:
            vehicle.make = make.strip() or None
            vehicle.model = model.strip() or None
            vehicle.year = int(year) if year.strip() else None
            vehicle.fuel_norm_per_100km = float(fuel_norm_per_100km) if fuel_norm_per_100km.strip() else None
            vehicle.maintenance_interval_km = int(maintenance_interval_km) if maintenance_interval_km.strip() else 60000
            vehicle.last_maintenance_odometer = int(last_maintenance_odometer) if last_maintenance_odometer.strip() else 0
            vehicle.trailer_plate = trailer_plate.upper().strip() or None
            session.add(vehicle)
            await session.commit()
    return RedirectResponse(url="/panel/drivers", status_code=302)


@router.get("/expenses", response_class=HTMLResponse)
async def list_expenses(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(Expense)
            .join(Trip, Expense.trip_id == Trip.id)
            .join(Driver, Trip.driver_id == Driver.id)
            .where(Expense.status == "PENDING")
            .order_by(Expense.created_at.desc())
            .options(
                selectinload(Expense.trip).selectinload(Trip.driver)
            )
        )
        expenses = result.scalars().all()
    return templates.TemplateResponse(request, "expenses.html",
        {"request": request, "expenses": expenses, "user": current_user},
    )


# --- EXCEL DOCUMENT GENERATOR ---
@router.get("/documents/generator", response_class=HTMLResponse)
async def docs_generator_get(request: Request, current_user: dict = Depends(get_current_user_from_cookie)):
    """Render the Excel Generator page for dispatchers."""
    async with async_session() as session:
        # Load trips to attach data to doc
        result = await session.execute(
            select(Trip).options(selectinload(Trip.driver), selectinload(Trip.vehicle)).order_by(Trip.id.desc())
        )
        trips = result.scalars().all()
    
    return templates.TemplateResponse(request, "documents_generator.html",
        {"request": request, "user": current_user, "trips": trips}
    )

import io
import re
from fastapi.responses import Response

@router.post("/documents/generate_excel")
async def docs_generate_excel_post(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie)
):
    """Generates an Excel file by replacing {{tags}} in the template with user payload."""
    form_data = await request.form()
    trip_id = form_data.get("trip_id")
    template_name = form_data.get("template_name", "cmr_template.xlsx")
    
    # Secure template path
    import os
    safe_name = re.sub(r'[^a-zA-Z0-9_\-\.]', '', template_name)
    template_path = os.path.join("src", "assets", "docs_templates", safe_name)
    
    if not os.path.exists(template_path):
        return templates.TemplateResponse(request, "documents_generator.html",
            {"request": request, "user": current_user, "trips": [], "error": f"Шаблон {safe_name} не найден. Обратитесь к админу."}
        )
        
    context_data = {}
    
    # 1. Autofill from Trip if selected
    if trip_id and str(trip_id).isdigit():
        async with async_session() as session:
            result = await session.execute(
                select(Trip).options(selectinload(Trip.driver), selectinload(Trip.vehicle)).where(Trip.id == int(trip_id))
            )
            trip = result.scalars().first()
            if trip:
                # Heuristic for route parsing "A -> B"
                r_from, r_to = "", ""
                if trip.route_details and "→" in trip.route_details:
                    parts = trip.route_details.split("→")
                    r_from, r_to = parts[0].strip(), parts[1].strip()
                elif trip.route_details and "->" in trip.route_details:
                    parts = trip.route_details.split("->")
                    r_from, r_to = parts[0].strip(), parts[1].strip()
                elif trip.route_details and "-" in trip.route_details:
                    parts = trip.route_details.split("-")
                    r_from, r_to = parts[0].strip(), parts[1].strip()
                else:
                    r_from = trip.route_details or ""

                context_data = {
                    "route_from": r_from,
                    "route_to": r_to,
                    "driver_name": trip.driver.name if trip.driver else "",
                    "truck_plate": trip.vehicle.license_plate if trip.vehicle else "",
                    "trailer_plate": trip.vehicle.license_plate if trip.vehicle else "", # Fallback to tractor plate
                    "client_name": trip.client_name or "",
                    "cargo_name": trip.load_info or "",
                    "cargo_weight": "",
                    "price": str(trip.price) if trip.price else "",
                }
                
    # 2. Override with manual form variables (data[xxx])
    for key, value in form_data.items():
        if key.startswith("data[") and key.endswith("]"):
            var_name = key[5:-1]
            if value and value.strip(): # if user explicitly typed something
                context_data[var_name] = value.strip()
                
    # 3. Process openpyxl
    try:
        import openpyxl
        wb = openpyxl.load_workbook(template_path)
        
        # We loop through all cells in all sheets and replace {{var_name}}
        pattern = re.compile(r'\{\{(.*?)\}\}')
        
        def replace_text(text, ctx):
            if not text: return text
            matches = pattern.finditer(text)
            for m in matches:
                tag = m.group(0)
                var = m.group(1).strip()
                val = ctx.get(var, "")
                text = text.replace(tag, val)
            return text

        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        original = cell.value
                        new_val = replace_text(original, context_data)
                        if new_val != original:
                            cell.value = new_val

        # Save to memory
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        headers = {
            "Content-Disposition": f'attachment; filename="Generated_{safe_name}"'
        }
        return Response(output.read(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)
        
    except Exception as e:
        return templates.TemplateResponse(request, "documents_generator.html",
            {"request": request, "user": current_user, "trips": [], "error": f"Ошибка сборки Excel: {str(e)}"}
        )


@router.post("/expenses/{expense_id}/approve", response_class=HTMLResponse)
async def approve_expense(
    request: Request,
    expense_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(Expense)
            .where(Expense.id == expense_id)
            .options(selectinload(Expense.trip).selectinload(Trip.driver))
        )
        expense = result.scalars().first()
        if not expense:
            from fastapi import HTTPException
            raise HTTPException(status_code=404, detail="Expense not found")
        expense.status = "APPROVED"
        expense.approved_by = current_user.get("sub", "dispatcher")
        expense.approved_at = datetime.utcnow()
        session.add(expense)
        await session.commit()
        await session.refresh(expense)
    await notify_driver_expense_decision(expense, approved=True)
    return templates.TemplateResponse(request, "expense_row.html",
        {"request": request, "expense": expense},
    )


@router.post("/expenses/{expense_id}/reject", response_class=HTMLResponse)
async def reject_expense(
    request: Request,
    expense_id: int,
    comment: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(Expense)
            .where(Expense.id == expense_id)
            .options(selectinload(Expense.trip).selectinload(Trip.driver))
        )
        expense = result.scalars().first()
        if not expense:
            from fastapi import HTTPException
            raise HTTPException(status_code=404, detail="Expense not found")
        expense.status = "REJECTED"
        expense.rejection_comment = comment or None
        expense.approved_by = current_user.get("sub", "dispatcher")
        expense.approved_at = datetime.utcnow()
        session.add(expense)
        await session.commit()
        await session.refresh(expense)
    await notify_driver_expense_decision(expense, approved=False, rejection_comment=comment)
    return templates.TemplateResponse(request, "expense_row.html",
        {"request": request, "expense": expense},
    )



@router.get("/requests", response_class=HTMLResponse)
async def list_requests(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        # pending
        res_pending = await session.execute(
            select(DriverRequest)
            .options(selectinload(DriverRequest.driver), selectinload(DriverRequest.trip))
            .where(DriverRequest.status == RequestStatus.PENDING)
            .order_by(DriverRequest.created_at.desc())
        )
        pending = res_pending.scalars().all()
        # history
        res_hist = await session.execute(
            select(DriverRequest)
            .options(selectinload(DriverRequest.driver), selectinload(DriverRequest.trip))
            .where(DriverRequest.status != RequestStatus.PENDING)
            .order_by(DriverRequest.resolved_at.desc())
            .limit(30)
        )
        history = res_hist.scalars().all()

    return templates.TemplateResponse(request, "requests.html", {
        "request": request, "pending": pending, "history": history,
        "user": current_user, "pending_requests_count": len(pending),
    })

@router.post("/requests/{req_id}/resolve")
async def resolve_request(
    req_id: int,
    action: str = Form(...), # APPROVE or REJECT
    comment: str = Form(""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        res = await session.execute(select(DriverRequest).where(DriverRequest.id == req_id))
        req = res.scalars().first()
        if req:
            req.status = RequestStatus.RESOLVED if action == 'APPROVE' else RequestStatus.REJECTED
            req.resolution_comment = comment.strip()
            req.resolved_at = func.now()
            await session.commit()
    return RedirectResponse("/panel/requests", status_code=302)


@router.get("/anomalies", response_class=HTMLResponse)
async def list_anomalies(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(FuelAnomaly)
            .where(FuelAnomaly.is_false_alarm == False)
            .order_by(FuelAnomaly.created_at.desc())
        )
        anomalies = result.scalars().all()
    return templates.TemplateResponse(request, "anomalies.html",
        {"request": request, "anomalies": anomalies, "user": current_user},
    )


@router.post("/anomalies/{anomaly_id}/dismiss", response_class=HTMLResponse)
async def dismiss_anomaly(
    request: Request,
    anomaly_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(FuelAnomaly).where(FuelAnomaly.id == anomaly_id)
        )
        anomaly = result.scalars().first()
        if anomaly:
            anomaly.is_false_alarm = True
            session.add(anomaly)
            await session.commit()
            await session.refresh(anomaly)
    # Return a replacement row showing "false alarm" state
    dismissed_row = f'<tr id="anomaly-{anomaly_id}"><td colspan="7"><em>Отмечена как ложная тревога</em></td></tr>'
    return HTMLResponse(content=dismissed_row)


@router.get("/trips", response_class=HTMLResponse)
async def list_trips(
    request: Request,
    driver_id: int | None = None,
    vehicle_id: int | None = None,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        stmt = (
            select(Trip)
            .options(selectinload(Trip.driver), selectinload(Trip.vehicle))
            .order_by(Trip.created_at.desc())
        )
        if driver_id is not None:
            stmt = stmt.where(Trip.driver_id == driver_id)
        if vehicle_id is not None:
            stmt = stmt.where(Trip.vehicle_id == vehicle_id)
        result = await session.execute(stmt)
        trips = result.scalars().all()

        drivers_result = await session.execute(
            select(Driver).order_by(Driver.name)
        )
        drivers = drivers_result.scalars().all()

        vehicles_result = await session.execute(
            select(Vehicle).order_by(Vehicle.license_plate)
        )
        vehicles = vehicles_result.scalars().all()

    return templates.TemplateResponse(request, "trips_list.html",
        {
            "request": request,
            "trips": trips,
            "drivers": drivers,
            "vehicles": vehicles,
            "filter_driver_id": driver_id,
            "filter_vehicle_id": vehicle_id,
            "user": current_user,
        },
    )


@router.get("/trips/create", response_class=HTMLResponse)
async def create_trip_form(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        drivers_result = await session.execute(
            select(Driver).order_by(Driver.name)
        )
        drivers = drivers_result.scalars().all()

        vehicles_result = await session.execute(
            select(Vehicle).order_by(Vehicle.license_plate)
        )
        vehicles = vehicles_result.scalars().all()

    return templates.TemplateResponse(request, "trips_create.html",
        {
            "request": request,
            "drivers": drivers,
            "vehicles": vehicles,
            "TripStatus": TripStatus,
            "user": current_user
        },
    )


@router.post("/trips/create")
async def create_trip_submit(
    request: Request,
    driver_id: int = Form(),
    vehicle_id: int = Form(),
    status: str = Form(default="PLANNED"),
    route_details: str = Form(default=""),
    load_info: str = Form(default=""),
    client_name: str = Form(default=""),
    planned_departure: str = Form(default=""),
    planned_arrival: str = Form(default=""),
    price: float = Form(default=None),
    price_currency: str = Form(default="KZT"),
    planned_distance_km: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    from datetime import datetime as dt

    def parse_dt(value: str):
        if not value:
            return None
        try:
            return dt.fromisoformat(value)
        except ValueError:
            return None

    def parse_int(value: str):
        if not value:
            return None
        try:
            return int(value)
        except ValueError:
            return None

    async with async_session() as session:
        # Check maintenance
        vehicle = (await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id).options(selectinload(Vehicle.trips)))).scalars().first()
        if vehicle:
            max_trip_odometer = max((t.end_odometer for t in vehicle.trips if t.end_odometer), default=0)
            current_odometer = max(max_trip_odometer, vehicle.last_maintenance_odometer or 0)
            km_since_maint = current_odometer - (vehicle.last_maintenance_odometer or 0)
            interval = vehicle.maintenance_interval_km or 60000
            
            if km_since_maint >= interval:
                drivers = (await session.execute(select(Driver).where(Driver.status == "ACTIVE").order_by(Driver.name))).scalars().all()
                vehicles = (await session.execute(select(Vehicle).order_by(Vehicle.license_plate))).scalars().all()
                return templates.TemplateResponse(request, "trips_create.html", {
                    "request": request, "user": current_user, "TripStatus": TripStatus,
                    "drivers": drivers, "vehicles": vehicles,
                    "error": f"Внимание! У выбранного борта ({vehicle.license_plate}) пройдено ТО или перепробег. Обновите ТО перед отправкой в рейс."
                }, status_code=400)

        trip = Trip(
            driver_id=driver_id,
            vehicle_id=vehicle_id,
            route_details=route_details or None,
            load_info=load_info or None,
            client_name=client_name or None,
            planned_departure=parse_dt(planned_departure),
            planned_arrival=parse_dt(planned_arrival),
            status=status,
            price=price,
            price_currency=price_currency,
            planned_distance_km=parse_int(planned_distance_km),
        )
        session.add(trip)
        await session.commit()
        await session.refresh(trip)

        # Sync vehicle status
        veh = (await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))).scalars().first()
        if veh:
            veh.status = VehicleStatus.ON_TRIP.value if status == TripStatus.ACTIVE.value else VehicleStatus.AVAILABLE.value
            await session.commit()

        # Eager-load driver for notification
        result = await session.execute(
            select(Trip)
            .where(Trip.id == trip.id)
            .options(selectinload(Trip.driver), selectinload(Trip.vehicle))
        )
        trip = result.scalars().first()

        # Load open repairs for this vehicle to include in notification
        repairs_result = await session.execute(
            select(VehicleRepair)
            .where(VehicleRepair.vehicle_id == vehicle_id)
            .where(VehicleRepair.status.in_([RepairStatus.OPEN, RepairStatus.IN_PROGRESS]))
            .order_by(VehicleRepair.created_at.desc())
        )
        open_repairs = repairs_result.scalars().all()

    await notify_driver_trip_assignment(trip, open_repairs=open_repairs)
    return RedirectResponse(url=f"/panel/trips/{trip.id}", status_code=302)


@router.get("/trips/{trip_id}", response_class=HTMLResponse)
async def trip_detail(
    request: Request,
    trip_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(Trip)
            .where(Trip.id == trip_id)
            .options(
                selectinload(Trip.driver),
                selectinload(Trip.vehicle),
                selectinload(Trip.expenses),
                selectinload(Trip.documents),
                selectinload(Trip.events),
            )
        )
        trip = result.scalars().first()
        if not trip:
            return RedirectResponse(url="/panel/trips", status_code=302)

        # Load location pings for this trip
        locations_result = await session.execute(
            select(LocationUpdate)
            .where(LocationUpdate.trip_id == trip_id)
            .order_by(LocationUpdate.created_at.desc())
        )
        locations = locations_result.scalars().all()

        # Load open/in-progress repairs for this vehicle
        repairs_result = await session.execute(
            select(VehicleRepair)
            .where(VehicleRepair.vehicle_id == trip.vehicle_id)
            .where(VehicleRepair.status.in_([RepairStatus.OPEN, RepairStatus.IN_PROGRESS]))
            .order_by(VehicleRepair.created_at.desc())
        )
        open_repairs = repairs_result.scalars().all()

        # For handover forms on active trips
        drivers_result = await session.execute(select(Driver).order_by(Driver.name))
        all_drivers = drivers_result.scalars().all()
        vehicles_result = await session.execute(select(Vehicle).order_by(Vehicle.license_plate))
        all_vehicles = vehicles_result.scalars().all()

    # Calculate total expenses in KZT
    total_kzt = sum(
        (e.amount_kzt if e.amount_kzt is not None else e.amount)
        for e in trip.expenses
    )

    return templates.TemplateResponse(request, "trip_detail.html",
        {
            "request": request,
            "trip": trip,
            "locations": locations,
            "total_kzt": total_kzt,
            "open_repairs": open_repairs,
            "drivers": all_drivers,
            "vehicles": all_vehicles,
            "user": current_user,
        },
    )


@router.get("/trips/{trip_id}/edit", response_class=HTMLResponse)
async def trip_edit_form(
    request: Request,
    trip_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(Trip).where(Trip.id == trip_id)
        )
        trip = result.scalars().first()
        if not trip:
            return RedirectResponse(url="/panel/trips", status_code=302)

        drivers_result = await session.execute(
            select(Driver).order_by(Driver.name)
        )
        drivers = drivers_result.scalars().all()

        vehicles_result = await session.execute(
            select(Vehicle).order_by(Vehicle.license_plate)
        )
        vehicles = vehicles_result.scalars().all()

    return templates.TemplateResponse(request, "trip_edit.html",
        {
            "request": request,
            "trip": trip,
            "drivers": drivers,
            "vehicles": vehicles,
            "TripStatus": TripStatus,
            "user": current_user,
        },
    )


@router.post("/trips/{trip_id}/edit")
async def trip_edit_submit(
    request: Request,
    trip_id: int,
    driver_id: int = Form(),
    vehicle_id: int = Form(),
    status: str = Form(),
    delivery_status: str = Form(default=""),
    route_details: str = Form(default=""),
    load_info: str = Form(default=""),
    client_name: str = Form(default=""),
    planned_departure: str = Form(default=""),
    planned_arrival: str = Form(default=""),
    start_odometer: str = Form(default=""),
    end_odometer: str = Form(default=""),
    price: float = Form(default=None),
    price_currency: str = Form(default="KZT"),
    planned_distance_km: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    from datetime import datetime as dt

    def parse_dt(value: str):
        if not value:
            return None
        try:
            return dt.fromisoformat(value)
        except ValueError:
            return None

    def parse_int(value: str):
        if not value:
            return None
        try:
            return int(value)
        except ValueError:
            return None

    async with async_session() as session:
        result = await session.execute(select(Trip).where(Trip.id == trip_id))
        trip = result.scalars().first()
        if not trip:
            return RedirectResponse(url="/panel/trips", status_code=302)

        trip.driver_id = driver_id
        trip.vehicle_id = vehicle_id
        trip.status = status
        trip.delivery_status = delivery_status.strip() or None
        trip.route_details = route_details.strip() or None
        trip.load_info = load_info.strip() or None
        trip.client_name = client_name.strip() or None
        trip.planned_departure = parse_dt(planned_departure)
        trip.planned_arrival = parse_dt(planned_arrival)
        trip.start_odometer = parse_int(start_odometer)
        trip.end_odometer = parse_int(end_odometer)
        trip.price = price
        trip.price_currency = price_currency
        trip.planned_distance_km = parse_int(planned_distance_km)

        session.add(trip)

        # Sync vehicle status based on trip status
        veh = (await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))).scalars().first()
        if veh and veh.status != VehicleStatus.MAINTENANCE.value:
            if status == TripStatus.ACTIVE.value:
                veh.status = VehicleStatus.ON_TRIP.value
            elif status in (TripStatus.COMPLETED.value, TripStatus.PLANNED.value):
                veh.status = VehicleStatus.AVAILABLE.value

        await session.commit()

    return RedirectResponse(url=f"/panel/trips/{trip.id}", status_code=302)


# --- FUEL ANTIFRAUD ---
@router.get("/antifraud", response_class=HTMLResponse)
async def antifraud_page(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    return templates.TemplateResponse(request, "antifraud.html", {
        "request": request, "user": current_user, "red_flags": None
    })

@router.post("/antifraud/run", response_class=HTMLResponse)
async def antifraud_run(
    request: Request,
    trans_file: UploadFile = File(...),
    manual_file: UploadFile = File(None),
    use_db: bool = Form(default=False),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    # 1. Read transaction file
    trans_content = (await trans_file.read()).decode("utf-8", errors="ignore")
    trans_data = FuelAntifraudService.parse_station_transactions(trans_content)
    
    manual_data = []
    
    # 2. Get manual data (either from uploaded file or DB)
    if use_db:
        async with async_session() as session:
            manual_data = await FuelAntifraudService.get_manual_data_from_db(session)
    
    if manual_file and manual_file.filename:
        manual_content = (await manual_file.read()).decode("utf-8", errors="ignore")
        manual_from_file = FuelAntifraudService.parse_manual_csv(manual_content)
        manual_data.extend(manual_from_file)
        
    if not manual_data and not use_db:
         return templates.TemplateResponse(request, "antifraud.html", {
            "request": request, "user": current_user, "error": "Необходимо загрузить ручной отчет или выбрать 'Сверку с БД'"
        })

    # 3. Run check
    red_flags = FuelAntifraudService.check_fraud(manual_data, trans_data)
    
    return templates.TemplateResponse(request, "antifraud.html", {
        "request": request, 
        "user": current_user, 
        "red_flags": red_flags,
        "processed_trans": len(trans_data),
        "processed_manual": len(manual_data)
    })


# --- PREVENTIVE MAINTENANCE ---
@router.get("/maintenance", response_class=HTMLResponse)
async def maintenance_dashboard(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    from src.services.maintenance import get_maintenance_dashboard_data
    async with async_session() as session:
        data = await get_maintenance_dashboard_data(session)
    return templates.TemplateResponse(request, "maintenance.html", {
        "request": request, "user": current_user, "vehicles": data
    })

@router.post("/maintenance/{vehicle_id}/reset")
async def maintenance_reset(
    vehicle_id: int,
    current_odo: int = Form(...),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
        if vehicle:
            vehicle.last_maintenance_odometer = current_odo
            session.add(vehicle)
            await session.commit()
    return RedirectResponse(url="/panel/maintenance", status_code=302)


# --- GEOFENCES ---
@router.get("/geofences", response_class=HTMLResponse)
async def geofences_list(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        res = await session.execute(select(Geofence).order_by(Geofence.created_at.desc()))
        fences = res.scalars().all()
    return templates.TemplateResponse(request, "geofences.html", {
        "request": request, "user": current_user, "geofences": fences
    })

@router.post("/geofences/add")
async def geofence_add(
    name: str = Form(...),
    latitude: float = Form(...),
    longitude: float = Form(...),
    radius: int = Form(500),
    event_type: str = Form(...),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        fence = Geofence(
            name=name.strip(),
            latitude=latitude,
            longitude=longitude,
            radius_meters=radius,
            event_type=event_type
        )
        session.add(fence)
        await session.commit()
    return RedirectResponse(url="/panel/geofences", status_code=302)

@router.post("/geofences/{fence_id}/delete")
async def geofence_delete(
    fence_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        res = await session.execute(select(Geofence).where(Geofence.id == fence_id))
        fence = res.scalars().first()
        if fence:
            await session.delete(fence)
            await session.commit()
    return RedirectResponse(url="/panel/geofences", status_code=302)


# --- ECO-DRIVING LEADERBOARD ---
@router.get("/eco-driving", response_class=HTMLResponse)
async def eco_driving_dashboard(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    from src.services.eco_driving import get_driver_leaderboard
    async with async_session() as session:
        leaderboard = await get_driver_leaderboard(session)
    return templates.TemplateResponse(request, "eco_driving.html", {
        "request": request, "user": current_user, "leaderboard": leaderboard
    })


@router.post("/trips/{trip_id}/delete")
async def trip_delete(
    trip_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Trip).where(Trip.id == trip_id))
        trip = result.scalars().first()
        if trip:
            # Clean up foreign keys that don't have cascade
            await session.execute(delete(FuelAnomaly).where(FuelAnomaly.trip_id == trip_id))
            await session.execute(sa_update(LocationUpdate).where(LocationUpdate.trip_id == trip_id).values(trip_id=None))
            await session.execute(sa_update(VehicleRepair).where(VehicleRepair.trip_id == trip_id).values(trip_id=None))
            await session.execute(sa_update(DriverRequest).where(DriverRequest.trip_id == trip_id).values(trip_id=None))

            await session.delete(trip)
            await session.commit()
    return RedirectResponse(url="/panel/trips", status_code=302)


@router.post("/vehicles/{vehicle_id}/trailer")
async def vehicle_trailer_update(
    vehicle_id: int,
    trailer_plate: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
        if vehicle:
            vehicle.trailer_plate = trailer_plate.strip().upper() or None
            await session.commit()
    return RedirectResponse(url="/panel/fleet", status_code=302)


@router.post("/trips/{trip_id}/change-driver")
async def trip_change_driver(
    trip_id: int,
    new_driver_id: int = Form(),
    note: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        trip_result = await session.execute(
            select(Trip).options(selectinload(Trip.driver)).where(Trip.id == trip_id)
        )
        trip = trip_result.scalars().first()
        if trip:
            old_name = trip.driver.name if trip.driver else str(trip.driver_id)
            new_drv_result = await session.execute(select(Driver).where(Driver.id == new_driver_id))
            new_driver = new_drv_result.scalars().first()
            new_name = new_driver.name if new_driver else str(new_driver_id)

            trip.driver_id = new_driver_id
            label = f"{old_name} → {new_name}"
            if note.strip():
                label += f" | {note.strip()}"
            session.add(TripEvent(
                trip_id=trip_id,
                event_type=TripEventType.DRIVER_CHANGE.value,
                location_name=label,
            ))
            await session.commit()
    return RedirectResponse(url=f"/panel/trips/{trip_id}", status_code=302)


@router.post("/trips/{trip_id}/change-vehicle")
async def trip_change_vehicle(
    trip_id: int,
    new_vehicle_id: int = Form(),
    note: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        trip_result = await session.execute(
            select(Trip).options(selectinload(Trip.vehicle)).where(Trip.id == trip_id)
        )
        trip = trip_result.scalars().first()
        if trip:
            old_plate = trip.vehicle.license_plate if trip.vehicle else str(trip.vehicle_id)

            # Free old vehicle
            old_veh_result = await session.execute(select(Vehicle).where(Vehicle.id == trip.vehicle_id))
            old_veh = old_veh_result.scalars().first()
            if old_veh and old_veh.status != VehicleStatus.MAINTENANCE.value:
                old_veh.status = VehicleStatus.AVAILABLE.value

            # Assign new vehicle
            new_veh_result = await session.execute(select(Vehicle).where(Vehicle.id == new_vehicle_id))
            new_veh = new_veh_result.scalars().first()
            new_plate = new_veh.license_plate if new_veh else str(new_vehicle_id)
            if new_veh and new_veh.status != VehicleStatus.MAINTENANCE.value:
                new_veh.status = VehicleStatus.ON_TRIP.value

            trip.vehicle_id = new_vehicle_id
            label = f"{old_plate} → {new_plate}"
            if note.strip():
                label += f" | {note.strip()}"
            session.add(TripEvent(
                trip_id=trip_id,
                event_type=TripEventType.VEHICLE_CHANGE.value,
                location_name=label,
            ))
            await session.commit()
    return RedirectResponse(url=f"/panel/trips/{trip_id}", status_code=302)


@router.get("/trips/{trip_id}/expenses/add", response_class=HTMLResponse)
async def add_trip_expense_form(
    request: Request,
    trip_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    from src.db.models import ExpenseType
    async with async_session() as session:
        result = await session.execute(select(Trip).where(Trip.id == trip_id))
        trip = result.scalars().first()
        if not trip:
            return RedirectResponse(url="/panel/trips", status_code=302)

    return templates.TemplateResponse(request, "expense_add.html",
        {
            "request": request,
            "trip": trip,
            "ExpenseType": ExpenseType,
            "user": current_user,
        },
    )


@router.post("/trips/{trip_id}/expenses/add")
async def add_trip_expense_submit(
    request: Request,
    trip_id: int,
    expense_type: str = Form(),
    amount: float = Form(),
    currency: str = Form(),
    status: str = Form(),
    comment: str = Form(default=""),
    liters: float = Form(default=None),
    odometer_reading: int = Form(default=None),
    station: str = Form(default=""),
    receipt_file: UploadFile = File(default=None),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    import uuid
    # Static FX rates to KZT
    FX_RATES: dict[str, float] = {
        "KZT": 1.0,
        "RUB": 5.2,
        "USD": 450.0,
        "EUR": 490.0,
        "TRY": 13.5,
        "CNY": 65.0,
    }
    fx_rate = FX_RATES.get(currency, 1.0)
    amount_kzt = round(amount * fx_rate, 2)

    receipt_url = None
    if receipt_file and receipt_file.filename:
        file_bytes = await receipt_file.read()
        if file_bytes:
            ext = receipt_file.filename.split('.')[-1] if '.' in receipt_file.filename else 'jpg'
            key = f"expenses/receipt_{uuid.uuid4()}.{ext}"
            content_type = receipt_file.content_type or "image/jpeg"
            receipt_url = await save_uploaded_file(file_bytes, key, content_type)

    async with async_session() as session:
        result = await session.execute(select(Trip).where(Trip.id == trip_id))
        trip = result.scalars().first()
        if not trip:
            return RedirectResponse(url="/panel/trips", status_code=302)

        expense = Expense(
            trip_id=trip.id,
            expense_type=expense_type,
            amount=amount,
            currency=currency,
            fx_rate=fx_rate,
            amount_kzt=amount_kzt,
            comment=comment.strip() or None,
            photo_r2_url=receipt_url,
            status=status,
        )
        if expense_type == "FUEL":
            expense.liters = liters
            expense.odometer_reading = odometer_reading
            expense.station = station.strip() or None

        if status == "APPROVED":
            expense.approved_by = current_user.get("sub", "dispatcher")
            expense.approved_at = datetime.utcnow()

        session.add(expense)
        await session.commit()

    return RedirectResponse(url=f"/panel/trips/{trip.id}", status_code=302)


@router.post("/trips/{trip_id}/documents")
async def add_trip_document(
    request: Request,
    trip_id: int,
    file_type: str = Form(...),
    doc_file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    import uuid
    async with async_session() as session:
        result = await session.execute(select(Trip).where(Trip.id == trip_id))
        trip = result.scalars().first()
        if not trip or not doc_file.filename:
            return RedirectResponse(url=f"/panel/trips/{trip_id}", status_code=302)

        file_bytes = await doc_file.read()
        ext = doc_file.filename.rsplit(".", 1)[-1].lower() if "." in doc_file.filename else "pdf"
        content_type = "application/pdf" if ext == "pdf" else "image/jpeg"
        key = f"trip_docs/{trip_id}/{uuid.uuid4()}.{ext}"
        r2_url = await save_uploaded_file(file_bytes, key, content_type)

        doc = TripDocument(
            trip_id=trip_id,
            file_type=file_type,
            r2_url=r2_url,
        )
        session.add(doc)
        await session.commit()
    return RedirectResponse(url=f"/panel/trips/{trip_id}", status_code=302)


@router.get("/fleet", response_class=HTMLResponse)
async def fleet_dashboard(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        vehicles_result = await session.execute(
            select(Vehicle)
            .options(
                selectinload(Vehicle.trips).selectinload(Trip.driver),
                selectinload(Vehicle.drivers),
                selectinload(Vehicle.repairs),
            )
            .order_by(Vehicle.license_plate)
        )
        vehicles = vehicles_result.scalars().all()

    vehicle_data = []
    for vehicle in vehicles:
        active_trip = next(
            (t for t in vehicle.trips if t.status == TripStatus.ACTIVE),
            None,
        )
        open_repairs = sum(
            1 for r in vehicle.repairs
            if r.status in (RepairStatus.OPEN, RepairStatus.IN_PROGRESS)
        )
        
        # Odometer calculation
        max_trip_odometer = max((t.end_odometer for t in vehicle.trips if t.end_odometer), default=0)
        current_odometer = max(max_trip_odometer, vehicle.last_maintenance_odometer or 0)
        
        interval = vehicle.maintenance_interval_km or 60000
        last_maint = vehicle.last_maintenance_odometer or 0
        km_since_maint = current_odometer - last_maint
        km_to_maint = interval - km_since_maint
        
        vehicle_data.append({
            "vehicle": vehicle, 
            "active_trip": active_trip, 
            "open_repairs": open_repairs,
            "km_to_maint": km_to_maint,
            "current_odometer": current_odometer
        })

    return templates.TemplateResponse(request, "fleet.html",
        {
            "request": request,
            "vehicle_data": vehicle_data,
            "VehicleStatus": VehicleStatus,
            "user": current_user,
        },
    )


async def _get_fleet_markers(session) -> list[dict]:
    """Return latest location for each driver with an active trip."""
    active_trips_result = await session.execute(
        select(Trip)
        .where(Trip.status == TripStatus.ACTIVE)
        .options(selectinload(Trip.driver), selectinload(Trip.vehicle))
    )
    active_trips = active_trips_result.scalars().all()

    if not active_trips:
        return []

    driver_ids = [t.driver_id for t in active_trips]

    # Latest LocationUpdate per driver
    subq = (
        select(
            LocationUpdate.driver_id,
            func.max(LocationUpdate.created_at).label("max_ts"),
        )
        .where(LocationUpdate.driver_id.in_(driver_ids))
        .group_by(LocationUpdate.driver_id)
        .subquery()
    )
    locs_result = await session.execute(
        select(LocationUpdate).join(
            subq,
            (LocationUpdate.driver_id == subq.c.driver_id)
            & (LocationUpdate.created_at == subq.c.max_ts),
        )
    )
    latest_locations = {loc.driver_id: loc for loc in locs_result.scalars().all()}

    markers = []
    for trip in active_trips:
        loc = latest_locations.get(trip.driver_id)
        if loc:
            markers.append(
                {
                    "lat": loc.latitude,
                    "lng": loc.longitude,
                    "driver_name": trip.driver.name or "—",
                    "plate": trip.vehicle.license_plate,
                    "updated_at": loc.created_at.strftime("%d.%m %H:%M")
                    if loc.created_at
                    else "—",
                }
            )
    return markers


@router.get("/map", response_class=HTMLResponse)
async def fleet_map(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        markers = await _get_fleet_markers(session)
    return templates.TemplateResponse(request, "map.html",
        {"request": request, "markers": markers, "user": current_user},
    )


@router.get("/map/markers")
async def fleet_map_markers(
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        markers = await _get_fleet_markers(session)
    return JSONResponse(content=markers)


# ── Vehicle Add ────────────────────────────────────────────────────────────

@router.get("/vehicles/add", response_class=HTMLResponse)
async def vehicle_add_form(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    return templates.TemplateResponse(request, "vehicle_add.html",
        {"request": request, "user": current_user},
    )


@router.post("/vehicles/add")
async def vehicle_add_submit(
    request: Request,
    license_plate: str = Form(),
    trailer_plate: str = Form(default=""),
    make: str = Form(default=""),
    model: str = Form(default=""),
    year: str = Form(default=""),
    fuel_norm_per_100km: str = Form(default=""),
    maintenance_interval_km: str = Form(default="60000"),
    last_maintenance_odometer: str = Form(default="0"),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    plate = license_plate.upper().strip()
    async with async_session() as session:
        existing = await session.execute(select(Vehicle).where(Vehicle.license_plate == plate))
        if existing.scalars().first():
            return templates.TemplateResponse(request, "vehicle_add.html",
                {"request": request, "user": current_user, "error": f"Борт {plate} уже существует"},
                status_code=400,
            )
        vehicle = Vehicle(
            license_plate=plate,
            trailer_plate=trailer_plate.upper().strip() or None,
            make=make.strip() or None,
            model=model.strip() or None,
            year=int(year) if year.strip() else None,
            fuel_norm_per_100km=float(fuel_norm_per_100km) if fuel_norm_per_100km.strip() else None,
            maintenance_interval_km=int(maintenance_interval_km) if maintenance_interval_km.strip() else 60000,
            last_maintenance_odometer=int(last_maintenance_odometer) if last_maintenance_odometer.strip() else 0,
        )
        session.add(vehicle)
        await session.commit()
    return RedirectResponse(url="/panel/fleet", status_code=302)


# ── Driver Add ─────────────────────────────────────────────────────────────

@router.get("/drivers/add", response_class=HTMLResponse)
async def driver_add_form(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        vehicles_result = await session.execute(select(Vehicle).order_by(Vehicle.license_plate))
        vehicles = vehicles_result.scalars().all()
    return templates.TemplateResponse(request, "driver_add.html",
        {"request": request, "vehicles": vehicles, "user": current_user},
    )


@router.post("/drivers/add")
async def driver_add_submit(
    request: Request,
    last_name: str = Form(default=""),
    first_name: str = Form(default=""),
    name: str = Form(default=""),
    phone: str = Form(default=""),
    iin: str = Form(default=""),
    vehicle_id: str = Form(default=""),
    password: str = Form(default=""),
    passport_file: UploadFile = File(default=None),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    # Support both split fields (new) and combined name (legacy)
    combined_name = f"{last_name.strip()} {first_name.strip()}".strip() if (last_name or first_name) else name.strip()

    passport_url = None
    if passport_file and passport_file.filename:
        file_bytes = await passport_file.read()
        if file_bytes:
            ext = passport_file.filename.rsplit(".", 1)[-1].lower()
            content_type = "application/pdf" if ext == "pdf" else "image/jpeg"
            key = f"passports/driver_{int(datetime.utcnow().timestamp())}_{passport_file.filename}"
            passport_url = await save_uploaded_file(file_bytes, key, content_type)

    vid = int(vehicle_id) if vehicle_id.strip() else None
    async with async_session() as session:
        driver = Driver(
            tg_id=None,
            name=combined_name,
            phone=phone.strip() or None,
            iin=iin.strip() or None,
            passport_r2_url=passport_url,
            vehicle_id=vid,
            password_hash=hash_password(password) if password else None,
            status="ACTIVE",
        )
        session.add(driver)
        await session.commit()
    return RedirectResponse(url="/panel/drivers", status_code=302)


# ── Driver Edit ────────────────────────────────────────────────────────────

@router.get("/drivers/{driver_id}/edit", response_class=HTMLResponse)
async def driver_edit_form(
    request: Request,
    driver_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Driver).where(Driver.id == driver_id))
        driver = result.scalars().first()
        if not driver:
            return RedirectResponse(url="/panel/drivers", status_code=302)
        vehicles_result = await session.execute(select(Vehicle).order_by(Vehicle.license_plate))
        vehicles = vehicles_result.scalars().all()
    return templates.TemplateResponse(request, "driver_edit.html",
        {"request": request, "driver": driver, "vehicles": vehicles, "user": current_user},
    )


@router.post("/drivers/{driver_id}/edit")
async def driver_edit_submit(
    request: Request,
    driver_id: int,
    last_name: str = Form(default=""),
    first_name: str = Form(default=""),
    name: str = Form(default=""),
    phone: str = Form(default=""),
    iin: str = Form(default=""),
    vehicle_id: str = Form(default=""),
    password: str = Form(default=""),
    passport_file: UploadFile = File(default=None),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    combined_name = f"{last_name.strip()} {first_name.strip()}".strip() if (last_name or first_name) else name.strip()

    async with async_session() as session:
        result = await session.execute(select(Driver).where(Driver.id == driver_id))
        driver = result.scalars().first()
        if not driver:
            return RedirectResponse(url="/panel/drivers", status_code=302)

        if combined_name:
            driver.name = combined_name
        if phone.strip():
            driver.phone = phone.strip()
        if iin.strip():
            driver.iin = iin.strip()
        if password.strip():
            driver.password_hash = hash_password(password)
        driver.vehicle_id = int(vehicle_id) if vehicle_id.strip() else None

        if passport_file and passport_file.filename:
            file_bytes = await passport_file.read()
            if file_bytes:
                ext = passport_file.filename.rsplit(".", 1)[-1].lower()
                content_type = "application/pdf" if ext == "pdf" else "image/jpeg"
                key = f"passports/driver_{driver_id}_{int(datetime.utcnow().timestamp())}.{ext}"
                driver.passport_r2_url = await save_uploaded_file(file_bytes, key, content_type)

        session.add(driver)
        await session.commit()
    return RedirectResponse(url="/panel/drivers", status_code=302)


# ── Repairs ────────────────────────────────────────────────────────────────

@router.get("/repairs", response_class=HTMLResponse)
async def list_repairs(
    request: Request,
    vehicle_id: int | None = None,
    status: str | None = None,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        stmt = (
            select(VehicleRepair)
            .options(
                selectinload(VehicleRepair.vehicle),
                selectinload(VehicleRepair.trip),
            )
            .order_by(VehicleRepair.created_at.desc())
        )
        if vehicle_id:
            stmt = stmt.where(VehicleRepair.vehicle_id == vehicle_id)
        if status:
            stmt = stmt.where(VehicleRepair.status == status)
        result = await session.execute(stmt)
        repairs = result.scalars().all()

        vehicles_result = await session.execute(select(Vehicle).order_by(Vehicle.license_plate))
        vehicles = vehicles_result.scalars().all()

    open_count = sum(1 for r in repairs if r.status == RepairStatus.OPEN)
    in_progress_count = sum(1 for r in repairs if r.status == RepairStatus.IN_PROGRESS)

    return templates.TemplateResponse(request, "repairs.html",
        {
            "request": request,
            "repairs": repairs,
            "vehicles": vehicles,
            "open_count": open_count,
            "in_progress_count": in_progress_count,
            "filter_vehicle_id": vehicle_id,
            "filter_status": status,
            "user": current_user,
            "RepairStatus": RepairStatus,
        },
    )


@router.post("/repairs")
async def create_repair(
    request: Request,
    vehicle_id: int = Form(),
    description: str = Form(),
    priority: str = Form(default="NORMAL"),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        repair = VehicleRepair(
            vehicle_id=vehicle_id,
            description=description.strip(),
            priority=priority,
            status=RepairStatus.OPEN,
            reported_by=current_user.get("sub", "dispatcher"),
        )
        session.add(repair)
        await session.commit()
    return RedirectResponse(url="/panel/repairs", status_code=302)


@router.post("/repairs/{repair_id}/start")
async def repair_start(
    repair_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(VehicleRepair).where(VehicleRepair.id == repair_id))
        repair = result.scalars().first()
        if repair and repair.status == RepairStatus.OPEN:
            repair.status = RepairStatus.IN_PROGRESS
            session.add(repair)
            await session.commit()
    return RedirectResponse(url="/panel/repairs", status_code=302)


@router.get("/repairs/{repair_id}/act", response_class=HTMLResponse)
async def repair_act(
    request: Request,
    repair_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    """Generate a printable breakdown act for a repair."""
    async with async_session() as session:
        result = await session.execute(
            select(VehicleRepair)
            .where(VehicleRepair.id == repair_id)
            .options(selectinload(VehicleRepair.vehicle))
        )
        repair = result.scalars().first()
        if not repair:
            raise HTTPException(status_code=404, detail="Repair not found")

    return templates.TemplateResponse(request, "repair_act.html",
        {"request": request, "repair": repair, "user": current_user},
    )


@router.post("/repairs/{repair_id}/resolve")
async def repair_resolve(
    repair_id: int,
    resolution_notes: str = Form(default=""),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(VehicleRepair).where(VehicleRepair.id == repair_id))
        repair = result.scalars().first()
        if repair and repair.status != RepairStatus.RESOLVED:
            repair.status = RepairStatus.RESOLVED
            repair.resolved_by = current_user.get("sub", "dispatcher")
            repair.resolution_notes = resolution_notes.strip() or None
            repair.resolved_at = datetime.utcnow()
            session.add(repair)
            await session.commit()
    return RedirectResponse(url="/panel/repairs", status_code=302)


@router.post("/repairs/{repair_id}/delete")
async def repair_delete(
    repair_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(VehicleRepair).where(VehicleRepair.id == repair_id))
        repair = result.scalars().first()
        if repair:
            await session.delete(repair)
            await session.commit()
    return RedirectResponse(url="/panel/repairs", status_code=302)


# ── Document expiry helper ────────────────────────────────────────────────────

def doc_expiry_status(expires_at) -> str:
    if not expires_at:
        return "ok"
    days = (expires_at - date_type.today()).days
    if days < 0:
        return "expired"
    if days <= 7:
        return "critical"
    if days <= 30:
        return "warning"
    return "ok"


DRIVER_DOC_LABELS = {
    "driver_license": "Водительское удостоверение",
    "passport": "Паспорт",
    "medical": "Медицинская справка",
    "contract": "Трудовой договор",
    "other": "Другой документ",
}

VEHICLE_DOC_LABELS = {
    "registration": "Техпаспорт",
    "insurance": "Страховой полис",
    "inspection": "Техосмотр",
    "customs": "Таможенные документы",
    "other": "Другой документ",
}


# ── Driver Documents ──────────────────────────────────────────────────────────

@router.get("/drivers/{driver_id}/documents", response_class=HTMLResponse)
async def driver_documents_get(
    request: Request,
    driver_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Driver).where(Driver.id == driver_id))
        driver = result.scalars().first()
        if not driver:
            raise HTTPException(status_code=404, detail="Driver not found")
        docs_result = await session.execute(
            select(DriverDocument).where(DriverDocument.driver_id == driver_id)
            .order_by(DriverDocument.uploaded_at.desc())
        )
        docs = docs_result.scalars().all()
    return templates.TemplateResponse(request, "driver_documents.html", {
        "request": request,
        "driver": driver,
        "docs": docs,
        "doc_expiry_status": doc_expiry_status,
        "DRIVER_DOC_LABELS": DRIVER_DOC_LABELS,
    })


@router.post("/drivers/{driver_id}/documents")
async def driver_documents_post(
    request: Request,
    driver_id: int,
    doc_type: str = Form(...),
    title: str = Form(""),
    expires_at: str = Form(""),
    doc_file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    import uuid
    async with async_session() as session:
        result = await session.execute(select(Driver).where(Driver.id == driver_id))
        driver = result.scalars().first()
        if not driver or not doc_file.filename:
            return RedirectResponse(url=f"/panel/drivers/{driver_id}/documents", status_code=302)

        file_bytes = await doc_file.read()
        ext = doc_file.filename.rsplit(".", 1)[-1].lower() if "." in doc_file.filename else "pdf"
        content_type = "application/pdf" if ext == "pdf" else "image/jpeg"
        key = f"driver_docs/{driver_id}/{uuid.uuid4()}.{ext}"
        r2_url = await save_uploaded_file(file_bytes, key, content_type)

        expires = None
        if expires_at.strip():
            try:
                from datetime import date as dt_date
                expires = dt_date.fromisoformat(expires_at.strip())
            except ValueError:
                pass

        doc = DriverDocument(
            driver_id=driver_id,
            doc_type=doc_type,
            title=title.strip() or None,
            r2_url=r2_url,
            expires_at=expires,
        )
        session.add(doc)
        await session.commit()
    return RedirectResponse(url=f"/panel/drivers/{driver_id}/documents", status_code=302)


@router.post("/drivers/{driver_id}/documents/{doc_id}/delete")
async def driver_document_delete(
    request: Request,
    driver_id: int,
    doc_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(DriverDocument).where(DriverDocument.id == doc_id, DriverDocument.driver_id == driver_id)
        )
        doc = result.scalars().first()
        if doc:
            await session.delete(doc)
            await session.commit()
    return RedirectResponse(url=f"/panel/drivers/{driver_id}/documents", status_code=302)


# ── Vehicle Documents ─────────────────────────────────────────────────────────

@router.get("/vehicles/{vehicle_id}/documents", response_class=HTMLResponse)
async def vehicle_documents_get(
    request: Request,
    vehicle_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
        if not vehicle:
            raise HTTPException(status_code=404, detail="Vehicle not found")
        docs_result = await session.execute(
            select(VehicleDocument).where(VehicleDocument.vehicle_id == vehicle_id)
            .order_by(VehicleDocument.uploaded_at.desc())
        )
        docs = docs_result.scalars().all()
    return templates.TemplateResponse(request, "vehicle_documents.html", {
        "request": request,
        "vehicle": vehicle,
        "docs": docs,
        "doc_expiry_status": doc_expiry_status,
        "VEHICLE_DOC_LABELS": VEHICLE_DOC_LABELS,
    })


@router.post("/vehicles/{vehicle_id}/documents")
async def vehicle_documents_post(
    request: Request,
    vehicle_id: int,
    doc_type: str = Form(...),
    title: str = Form(""),
    expires_at: str = Form(""),
    doc_file: UploadFile = File(...),
    current_user: dict = Depends(get_current_user_from_cookie),
):
    import uuid
    async with async_session() as session:
        result = await session.execute(select(Vehicle).where(Vehicle.id == vehicle_id))
        vehicle = result.scalars().first()
        if not vehicle or not doc_file.filename:
            return RedirectResponse(url=f"/panel/vehicles/{vehicle_id}/documents", status_code=302)

        file_bytes = await doc_file.read()
        ext = doc_file.filename.rsplit(".", 1)[-1].lower() if "." in doc_file.filename else "pdf"
        content_type = "application/pdf" if ext == "pdf" else "image/jpeg"
        key = f"vehicle_docs/{vehicle_id}/{uuid.uuid4()}.{ext}"
        r2_url = await save_uploaded_file(file_bytes, key, content_type)

        expires = None
        if expires_at.strip():
            try:
                from datetime import date as dt_date
                expires = dt_date.fromisoformat(expires_at.strip())
            except ValueError:
                pass

        doc = VehicleDocument(
            vehicle_id=vehicle_id,
            doc_type=doc_type,
            title=title.strip() or None,
            r2_url=r2_url,
            expires_at=expires,
        )
        session.add(doc)
        await session.commit()
    return RedirectResponse(url=f"/panel/vehicles/{vehicle_id}/documents", status_code=302)


@router.post("/vehicles/{vehicle_id}/documents/{doc_id}/delete")
async def vehicle_document_delete(
    request: Request,
    vehicle_id: int,
    doc_id: int,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    async with async_session() as session:
        result = await session.execute(
            select(VehicleDocument).where(VehicleDocument.id == doc_id, VehicleDocument.vehicle_id == vehicle_id)
        )
        doc = result.scalars().first()
        if doc:
            await session.delete(doc)
            await session.commit()
    return RedirectResponse(url=f"/panel/vehicles/{vehicle_id}/documents", status_code=302)


# ── Expiring Documents Dashboard ──────────────────────────────────────────────

@router.get("/documents/expiring", response_class=HTMLResponse)
async def documents_expiring(
    request: Request,
    current_user: dict = Depends(get_current_user_from_cookie),
):
    from datetime import timedelta
    today = date_type.today()
    deadline = today + timedelta(days=30)

    async with async_session() as session:
        driver_docs_res = await session.execute(
            select(DriverDocument, Driver)
            .join(Driver, DriverDocument.driver_id == Driver.id)
            .where(DriverDocument.expires_at != None)
            .where(DriverDocument.expires_at <= deadline)
            .order_by(DriverDocument.expires_at)
        )
        driver_docs = driver_docs_res.all()

        vehicle_docs_res = await session.execute(
            select(VehicleDocument, Vehicle)
            .join(Vehicle, VehicleDocument.vehicle_id == Vehicle.id)
            .where(VehicleDocument.expires_at != None)
            .where(VehicleDocument.expires_at <= deadline)
            .order_by(VehicleDocument.expires_at)
        )
        vehicle_docs = vehicle_docs_res.all()

    return templates.TemplateResponse(request, "documents_expiring.html", {
        "request": request,
        "driver_docs": driver_docs,
        "vehicle_docs": vehicle_docs,
        "doc_expiry_status": doc_expiry_status,
        "DRIVER_DOC_LABELS": DRIVER_DOC_LABELS,
        "VEHICLE_DOC_LABELS": VEHICLE_DOC_LABELS,
        "today": today,
    })
